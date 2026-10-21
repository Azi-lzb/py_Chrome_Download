import unittest
from io import StringIO
from pathlib import Path
from unittest.mock import patch
from unittest.mock import Mock

import main


class MatrixRecoveryTests(unittest.TestCase):
    def test_recovery_mode_helper_distinguishes_plain_and_recover(self):
        self.assertFalse(main._matrix_recovery_enabled("plain"))
        self.assertFalse(main._matrix_recovery_enabled(""))
        self.assertTrue(main._matrix_recovery_enabled("recover"))

        with self.assertRaises(ValueError):
            main._matrix_recovery_enabled("unknown")

    def test_field_lookup_errors_are_recoverable_only_for_missing_report_inputs(self):
        self.assertTrue(main._is_recoverable_field_lookup_error("input_not_found"))
        self.assertTrue(main._is_recoverable_field_lookup_error("reportBox_not_found"))
        self.assertTrue(main._is_recoverable_field_lookup_error("INPUT_NOT_FOUND"))

        self.assertFalse(main._is_recoverable_field_lookup_error("checkbox_dropdown_not_match"))
        self.assertFalse(main._is_recoverable_field_lookup_error(""))

    def test_download_result_requires_confirmed_new_file(self):
        self.assertTrue(main._download_result_ok({"ok": True, "files": [{"path": "a.xlsx"}]}))

        self.assertFalse(main._download_result_ok({"ok": False, "timeout": True, "files": []}))
        self.assertFalse(main._download_result_ok({"ok": True, "files": []}))
        self.assertFalse(main._download_result_ok(None))

    def test_query_timeout_is_stalled_only_while_loading(self):
        self.assertTrue(main._query_wait_indicates_stall({"timeout": True, "saw_loading": True}))
        self.assertFalse(main._query_wait_indicates_stall({"timeout": True, "saw_loading": False}))
        self.assertFalse(main._query_wait_indicates_stall({"timeout": True, "saw_loading": True, "loading": False}))
        self.assertFalse(main._query_wait_indicates_stall({"timeout": False, "saw_loading": True}))
        self.assertFalse(main._query_wait_indicates_stall(None))

    def test_matrix_field_skips_only_when_current_page_value_matches(self):
        self.assertTrue(main._should_skip_existing_matrix_field("广东省", "广东省"))
        self.assertFalse(main._should_skip_existing_matrix_field("", "广东省"))
        self.assertFalse(main._should_skip_existing_matrix_field("惠州市", "广东省"))
        self.assertTrue(main._should_skip_existing_matrix_field("", "", want_clear=True))

    def test_configured_field_settle_delay_applies_to_sendkey_and_dropdown(self):
        self.assertEqual(main._field_settle_delay("sendkey", "sendkey", 2), 2.0)
        self.assertEqual(main._field_settle_delay("checkbox_dropdown", "checkbox_dropdown", 2), 2.0)
        self.assertEqual(main._field_settle_delay("sendkey", "sendkey", 0), 0.35)

    def test_reset_config_only_appends_missing_keys_without_overwriting_values(self):
        from openpyxl import Workbook, load_workbook

        artifact_dir = Path("output") / "test_artifacts"
        artifact_dir.mkdir(parents=True, exist_ok=True)
        config_path = artifact_dir / "reset_config_preserve.xlsx"
        wb = Workbook()
        ws = wb.active
        ws.title = "org"
        ws.append(["key", "value", "note"])
        ws.append(["task.query_wait_seconds", "150", "keep this value"])
        ws.append(["custom.key", "custom value", "keep this row"])
        wb.save(config_path)

        with patch.object(main, "EXCEL_CONFIG_FILE", config_path):
            main.reset_config_template()

        result = load_workbook(config_path, data_only=True)
        org = result["org"]
        values = {
            str(org.cell(row, 1).value or "").strip(): org.cell(row, 2).value
            for row in range(1, org.max_row + 1)
        }
        self.assertEqual(values["task.query_wait_seconds"], "150")
        self.assertEqual(values["custom.key"], "custom value")
        self.assertIn("task.field_settle_seconds", values)

    def test_recoverable_field_lookup_failure_should_abort_batch_after_retries(self):
        self.assertTrue(main._should_abort_after_field_lookup_retries("input_not_found", True))
        self.assertTrue(main._should_abort_after_field_lookup_retries("reportBox_not_found", True))

        self.assertFalse(main._should_abort_after_field_lookup_retries("input_not_found", False))
        self.assertFalse(main._should_abort_after_field_lookup_retries("checkbox_dropdown_not_match", True))

    def test_recovery_step_requires_optional_field_validation(self):
        self.assertTrue(main._recovery_step_passed(True, True))
        self.assertFalse(main._recovery_step_passed(True, False))
        self.assertFalse(main._recovery_step_passed(False, True))

    def test_login_state_treats_portal_entry_as_logged_in(self):
        driver = Mock()
        driver.execute_script.return_value = {
            "expired": False,
            "on_login_page": False,
            "has_logged_in_ui": True,
            "has_portal_entry": True,
            "has_report_tree": False,
        }
        with patch.object(main, "_driver", driver):
            state = main._detect_login_state()

        self.assertTrue(state["logged_in"])
        self.assertFalse(state["on_login_page"])

    def test_login_state_detects_visible_login_form(self):
        driver = Mock()
        driver.execute_script.return_value = {
            "expired": False,
            "on_login_page": True,
            "has_logged_in_ui": False,
            "has_portal_entry": False,
            "has_report_tree": False,
        }
        with patch.object(main, "_driver", driver):
            state = main._detect_login_state()

        self.assertFalse(state["logged_in"])
        self.assertTrue(state["on_login_page"])

    def test_full_recovery_uses_login_portal_tree_and_field_validation(self):
        with (
            patch.object(main, "_ensure_logged_in", return_value=True) as login,
            patch.object(main, "_navigate_to_report_page", return_value=True) as portal,
            patch.object(main, "_navigate_tree_to_report", return_value=True) as tree,
            patch.object(main, "_switch_to_report_frame", return_value=True) as frame,
            patch.object(main, "_report_field_available", return_value=True) as field,
            patch.object(main, "_save_iframe_src") as save,
            patch.object(main, "_wait_for_recovery_step") as wait_step,
        ):
            ok = main._full_recover_report(
                "org",
                {"tree_path": "A > B"},
                validate_label="数据日期:",
                expand_tree_on_retry=False,
            )

        self.assertTrue(ok)
        login.assert_called_once()
        portal.assert_called_once()
        tree.assert_called_once_with("A > B")
        frame.assert_called_once_with("org", {"tree_path": "A > B"})
        field.assert_called_once_with("数据日期:")
        save.assert_called_once()
        self.assertEqual(
            [call.args[0] for call in wait_step.call_args_list],
            ["登录状态确认", "报表服务中心导航", "tree_path 导航", "报表 iframe 切换"],
        )

    def test_full_recovery_can_restart_browser_before_login(self):
        with (
            patch.object(main, "_restart_browser_session", return_value=True) as restart,
            patch.object(main, "auto_login") as auto_login,
            patch.object(main, "_ensure_logged_in", return_value=True) as login,
            patch.object(main, "_navigate_to_report_page", return_value=True),
            patch.object(main, "_navigate_tree_to_report", return_value=True),
            patch.object(main, "_switch_to_report_frame", return_value=True),
            patch.object(main, "_report_field_available", return_value=True),
            patch.object(main, "_save_iframe_src"),
            patch.object(main, "_wait_for_recovery_step"),
        ):
            ok = main._full_recover_report(
                "org",
                {"tree_path": "A > B"},
                validate_label="数据日期:",
                restart_browser=True,
            )

        self.assertTrue(ok)
        restart.assert_called_once()
        auto_login.assert_called_once()
        login.assert_called_once()

    def test_full_recovery_expands_tree_and_retries_when_tree_path_fails(self):
        with (
            patch.object(main, "_ensure_logged_in", return_value=True),
            patch.object(main, "_navigate_to_report_page", return_value=True),
            patch.object(main, "_navigate_tree_to_report", side_effect=[False, True]) as tree,
            patch.object(main, "_expand_all_report_tree", return_value=True) as expand,
            patch.object(main, "_click_tree_path_leaf", return_value=True) as click_leaf,
            patch.object(main, "_switch_to_report_frame", return_value=True),
            patch.object(main, "_report_field_available", return_value=True),
            patch.object(main, "_save_iframe_src"),
            patch.object(main, "_wait_for_recovery_step"),
        ):
            ok = main._full_recover_report(
                "org",
                {"tree_path": "A > B"},
                validate_label="数据日期:",
                expand_tree_on_retry=True,
            )

        self.assertTrue(ok)
        self.assertEqual(tree.call_count, 2)
        expand.assert_called_once()
        click_leaf.assert_called_once_with("A > B")

    def test_item_exception_recovery_logs_and_continues_when_recovered(self):
        tmp = Path("output") / "test_artifacts"
        tmp.mkdir(parents=True, exist_ok=True)
        log_file = tmp / "matrix_recovery_log.jsonl"
        if log_file.exists():
            log_file.unlink()
        summaries = []

        with patch.object(main, "_recover_report", return_value=True) as recover:
            recovered = main._handle_matrix_item_exception(
                RuntimeError("iframe crashed"),
                scene="org",
                plan={"tree_path": "A > B"},
                seq=4,
                log_file=log_file,
                record_item_summary=lambda seq, started, ok, status: summaries.append((seq, ok, status)),
                item_started=123.0,
            )

        self.assertTrue(recovered)
        recover.assert_called_once_with("org", {"tree_path": "A > B"})
        self.assertEqual(summaries, [(4, False, "item_exception_recovered")])
        self.assertIn('"type": "step_error"', log_file.read_text(encoding="utf-8"))
        self.assertIn("iframe crashed", log_file.read_text(encoding="utf-8"))

    def test_menu_promotes_recovery_downloads_to_formal_entries(self):
        buf = StringIO()
        with patch("sys.stdout", buf), patch.object(main, "browser_status_line", return_value="test"):
            main.print_menu()

        menu = buf.getvalue()
        self.assertIn("3.  地区汇总报表查询导出", menu)
        self.assertIn("4.  单家机构报表查询导出", menu)
        self.assertIn("5.  注册当地法人机构报表查询导出", menu)
        self.assertNotIn("county，强恢复兜底", menu)
        self.assertNotIn("org，强恢复兜底", menu)
        self.assertNotIn("county2，强恢复兜底", menu)
        self.assertNotIn("t.  测试软刷新", menu)
        self.assertNotIn("t2. 测试门户", menu)
        self.assertNotIn("t3. 测试硬刷新", menu)
        self.assertNotIn("t31. 测试3增强恢复版（county）", menu)
        self.assertNotIn("t41. 测试4增强恢复版（org）", menu)
        self.assertNotIn("t51. 测试5增强恢复版（county2）", menu)
        self.assertNotIn("3p", menu)
        self.assertNotIn("4p", menu)
        self.assertNotIn("5p", menu)

    def test_formal_download_entries_use_recovery_mode(self):
        cases = [
            ("3", ("county", "recover")),
            ("4", ("org", "recover")),
            ("5", ("county2", "recover")),
        ]

        for choice, expected in cases:
            with self.subTest(choice=choice):
                calls = []

                with (
                    patch("builtins.input", side_effect=[choice, "0"]),
                    patch.object(main, "print_menu"),
                    patch.object(main, "run_matrix_export", side_effect=lambda scene, recovery_mode='plain': calls.append((scene, recovery_mode))),
                ):
                    main.main()

                self.assertEqual(calls, [expected])

    def test_plain_download_aliases_are_hidden_and_use_plain_mode(self):
        cases = [
            ("3p", ("county", "plain")),
            ("4p", ("org", "plain")),
            ("5p", ("county2", "plain")),
        ]

        for choice, expected in cases:
            with self.subTest(choice=choice):
                calls = []

                with (
                    patch("builtins.input", side_effect=[choice, "0"]),
                    patch.object(main, "print_menu"),
                    patch.object(main, "run_matrix_export", side_effect=lambda scene, recovery_mode='plain': calls.append((scene, recovery_mode))),
                ):
                    main.main()

                self.assertEqual(calls, [expected])

    def test_old_recovery_test_aliases_are_removed(self):
        for choice in ("t31", "t41", "t51"):
            with self.subTest(choice=choice):
                with (
                    patch("builtins.input", side_effect=[choice, "0"]),
                    patch.object(main, "print_menu"),
                    patch.object(main, "run_matrix_export") as export,
                ):
                    main.main()

                export.assert_not_called()


if __name__ == "__main__":
    unittest.main()
